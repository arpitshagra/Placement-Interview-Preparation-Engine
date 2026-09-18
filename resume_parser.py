"""
resume_parser.py — Parse PDF and DOCX resumes to extract raw text.
Supports:
  - PDF  via PyMuPDF  (pip install pymupdf)
  - DOCX via python-docx (pip install python-docx)
"""

import os


def parse_pdf(filepath: str) -> str:
    """Extract raw text from a PDF file using PyMuPDF (fitz)."""
    import fitz  # PyMuPDF
    text_parts = []
    with fitz.open(filepath) as doc:
        for page in doc:
            text_parts.append(page.get_text("text"))
    return "\n".join(text_parts).strip()


def parse_docx(filepath: str) -> str:
    """Extract raw text from a DOCX file using python-docx."""
    from docx import Document
    doc = Document(filepath)
    parts = []

    # Paragraph text
    for para in doc.paragraphs:
        t = para.text.strip()
        if t:
            parts.append(t)

    # Table cell text
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                t = cell.text.strip()
                if t and t not in parts:
                    parts.append(t)

    return "\n".join(parts).strip()


def parse_resume(filepath: str, file_type: str) -> str:
    """
    Dispatch to the correct parser based on file_type ('pdf' or 'docx').
    Returns the full raw text of the resume.
    """
    ft = file_type.lower().strip(".")
    if ft == "pdf":
        return parse_pdf(filepath)
    elif ft == "docx":
        return parse_docx(filepath)
    else:
        raise ValueError(f"Unsupported resume file type: '{file_type}'. Use 'pdf' or 'docx'.")
